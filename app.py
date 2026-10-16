import os
from html import escape

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

# ---------------------------------------------------------
# WMS RACK - thao tac truc tiep tren so do
# Dat file Python cung thu muc voi RACK.xlsx
# ---------------------------------------------------------
st.set_page_config(
    page_title="WMS Rack Management",
    page_icon="📦",
    layout="wide",
    initial_sidebar_state="expanded",
)

EXCEL_PATH = "RACK.xlsx"

st.markdown(
    """
    <style>
    .stApp { background: #0b0f19; color: #e5e7eb; }
    section[data-testid="stSidebar"] { background: #111827; }
    div[data-testid="stButton"] > button {
        min-height: 42px; width: 100%; padding: 3px 2px;
        border-radius: 7px; font-size: 11px; font-weight: 700;
        white-space: normal; overflow-wrap: anywhere;
        border: 1px solid #334155; background: #1e293b; color: #e2e8f0;
    }
    div[data-testid="stButton"] > button:hover {
        border-color: #38bdf8; color: #fff; background: #26364d;
    }
    .kpi {
        padding: 14px 16px; border: 1px solid #334155; border-radius: 10px;
        background: linear-gradient(135deg,#1e293b,#0f172a);
    }
    .kpi-label { color:#94a3b8; font-size:11px; text-transform:uppercase; }
    .kpi-value { color:#38bdf8; font-size:23px; font-weight:800; }
    .hint { color:#94a3b8; font-size:13px; }
    </style>
    """,
    unsafe_allow_html=True,
)

if not os.path.exists(EXCEL_PATH):
    st.error(f"Không tìm thấy **{EXCEL_PATH}**. Hãy đặt file Excel cùng thư mục với file Python.")
    st.stop()


def sheet_names():
    return pd.ExcelFile(EXCEL_PATH).sheet_names


def read_sheet(sheet_name):
    return pd.read_excel(EXCEL_PATH, sheet_name=sheet_name, header=None).fillna("").astype(str)


def write_cell(sheet_name, row, col, value):
    """Ghi đúng một ô bằng openpyxl để giữ định dạng workbook tốt hơn."""
    wb = load_workbook(EXCEL_PATH)
    ws = wb[sheet_name]
    ws.cell(row=row + 1, column=col + 1).value = value
    wb.save(EXCEL_PATH)
    wb.close()
    st.cache_data.clear()


@st.cache_data
def cached_sheet(sheet_name):
    return read_sheet(sheet_name)


def is_label(value):
    v = str(value).strip().lower()
    return not v or "tang" in v or "khu" in v


def position_text(row, col):
    return f"Dòng {row + 1}, cột {col + 1}"


sheets = sheet_names()
if not sheets:
    st.error("File Excel không có sheet.")
    st.stop()

if "current_sheet" not in st.session_state or st.session_state.current_sheet not in sheets:
    st.session_state.current_sheet = sheets[0]
if "selected_pos" not in st.session_state:
    st.session_state.selected_pos = None
if "move_mode" not in st.session_state:
    st.session_state.move_mode = False
if "move_source" not in st.session_state:
    st.session_state.move_source = None
if "notice" not in st.session_state:
    st.session_state.notice = None

# Sidebar controls
st.sidebar.markdown("## ⚡ WMS RACK")
st.sidebar.caption("Quản lý vị trí thùng hàng trực quan")
selected_sheet = st.sidebar.selectbox(
    "Chọn Rack / Sheet",
    sheets,
    index=sheets.index(st.session_state.current_sheet),
)
if selected_sheet != st.session_state.current_sheet:
    st.session_state.current_sheet = selected_sheet
    st.session_state.selected_pos = None
    st.session_state.move_mode = False
    st.session_state.move_source = None
    st.rerun()

df = cached_sheet(st.session_state.current_sheet)
search = st.sidebar.text_input("🔎 Tìm mã thùng", placeholder="Nhập một phần mã, ví dụ ML-138").strip()

if search:
    matches = []
    for s in sheets:
        d = cached_sheet(s)
        for r in range(d.shape[0]):
            for c in range(d.shape[1]):
                val = str(d.iat[r, c]).strip()
                if search.lower() in val.lower() and not is_label(val):
                    matches.append((s, r, c, val))
    st.sidebar.markdown(f"**Tìm thấy: {len(matches)} vị trí**")
    if matches:
        choices = [f"{s} · {val} · Dòng {r+1}, cột {c+1}" for s, r, c, val in matches]
        choice = st.sidebar.selectbox("Chọn kết quả để định vị", choices)
        if st.sidebar.button("📍 Đi tới vị trí", use_container_width=True):
            s, r, c, val = matches[choices.index(choice)]
            st.session_state.current_sheet = s
            st.session_state.selected_pos = (r, c)
            st.session_state.move_mode = False
            st.session_state.move_source = None
            st.rerun()

st.sidebar.divider()
st.sidebar.markdown("### 🛠️ Thao tác")
action = st.sidebar.radio(
    "Chế độ",
    ["Chọn / xem vị trí", "Thêm / sửa thùng", "Di chuyển thùng", "Xóa / xuất hàng"],
    index=0,
)
st.sidebar.caption("Cách dùng: bấm trực tiếp vào ô trên sơ đồ để chọn vị trí.")

# Notice from previous action
if st.session_state.notice:
    st.info(st.session_state.notice)
    st.session_state.notice = None

# Header and KPIs
total_cells = int(df.shape[0] * df.shape[1])
occupied_cells = sum(
    1 for r in range(df.shape[0]) for c in range(df.shape[1])
    if not is_label(str(df.iat[r, c]).strip())
)
rate = round(occupied_cells / total_cells * 100, 1) if total_cells else 0

st.title(f"📦 Sơ đồ Rack: {st.session_state.current_sheet}")
st.caption(f"{df.shape[0]} dòng × {df.shape[1]} cột · Bấm vào ô để chọn trực tiếp")
k1, k2, k3 = st.columns(3)
with k1:
    st.markdown(f'<div class="kpi"><div class="kpi-label">Tổng số ô</div><div class="kpi-value">{total_cells}</div></div>', unsafe_allow_html=True)
with k2:
    st.markdown(f'<div class="kpi"><div class="kpi-label">Ô đang chứa hàng</div><div class="kpi-value">{occupied_cells}</div></div>', unsafe_allow_html=True)
with k3:
    st.markdown(f'<div class="kpi"><div class="kpi-label">Tỷ lệ lấp đầy</div><div class="kpi-value">{rate}%</div></div>', unsafe_allow_html=True)

st.write("")

# Current selection panel
selected = st.session_state.selected_pos
if selected and selected[0] < df.shape[0] and selected[1] < df.shape[1]:
    selected_value = str(df.iat[selected[0], selected[1]]).strip()
    st.markdown(
        f"**Ô đang chọn:** `{selected_value or '(trống)'}` · "
        f"{position_text(*selected)} · Sheet `{st.session_state.current_sheet}`"
    )
else:
    selected_value = ""
    st.markdown('<p class="hint">Chưa chọn ô nào. Hãy bấm vào một ô trong sơ đồ bên dưới.</p>', unsafe_allow_html=True)

if action == "Di chuyển thùng":
    if not st.session_state.move_mode:
        if selected and not is_label(selected_value):
            if st.button("🚚 Dùng ô đang chọn làm VỊ TRÍ NGUỒN"):
                st.session_state.move_source = (st.session_state.current_sheet, selected[0], selected[1])
                st.session_state.move_mode = True
                st.session_state.notice = f"Đã chọn nguồn: {selected_value}. Chuyển sang rack đích (nếu cần), rồi bấm ô đích trên sơ đồ."
                st.rerun()
        else:
            st.warning("Hãy bấm chọn ô đang có thùng hàng trước.")
    else:
        src_sheet, src_r, src_c = st.session_state.move_source
        src_df = cached_sheet(src_sheet)
        src_val = str(src_df.iat[src_r, src_c]).strip()
        st.info(f"Đang di chuyển **{src_val}** từ {src_sheet} — {position_text(src_r, src_c)}. Bấm một ô đích trên sơ đồ.")
        target_is_valid = selected is not None and (st.session_state.current_sheet, selected[0], selected[1]) != (src_sheet, src_r, src_c)
        if target_is_valid:
            target_val = str(df.iat[selected[0], selected[1]]).strip()
            a, b = st.columns([1, 1])
            with a:
                if st.button(f"✅ Chuyển đến ô này ({target_val or 'trống'})", use_container_width=True):
                    if is_label(target_val):
                        write_cell(src_sheet, src_r, src_c, "")
                        write_cell(st.session_state.current_sheet, selected[0], selected[1], src_val)
                        st.session_state.move_mode = False
                        st.session_state.move_source = None
                        st.session_state.notice = f"Đã chuyển {src_val} đến {st.session_state.current_sheet}, {position_text(*selected)}."
                        st.rerun()
                    else:
                        st.error("Ô đích đã có hàng. Hãy chọn ô trống để tránh ghi đè.")
            with b:
                if st.button("Hủy di chuyển", use_container_width=True):
                    st.session_state.move_mode = False
                    st.session_state.move_source = None
                    st.rerun()
        elif st.button("Hủy di chuyển"):
            st.session_state.move_mode = False
            st.session_state.move_source = None
            st.rerun()

elif action == "Thêm / sửa thùng":
    if selected:
        new_value = st.text_input("Mã thùng / nội dung", value=selected_value, key=f"edit_{st.session_state.current_sheet}_{selected[0]}_{selected[1]}")
        if st.button("💾 Lưu vào ô đang chọn", type="primary"):
            if not is_label(selected_value) or not is_label(new_value):
                write_cell(st.session_state.current_sheet, selected[0], selected[1], new_value.strip())
                st.session_state.notice = f"Đã cập nhật {position_text(*selected)}."
                st.rerun()
            else:
                st.warning("Không thể ghi đè ô nhãn tầng/khu hoặc ô trống bằng nội dung rỗng.")
    else:
        st.info("Bấm một ô trên sơ đồ trước, sau đó nhập mã thùng để thêm hoặc sửa.")

elif action == "Xóa / xuất hàng":
    if selected and not is_label(selected_value):
        st.warning(f"Sẽ xóa mã **{selected_value}** tại {position_text(*selected)}.")
        confirm = st.checkbox("Tôi xác nhận muốn xóa / xuất thùng này", key="confirm_delete")
        if st.button("🗑️ Xóa thùng khỏi vị trí", type="primary", disabled=not confirm):
            write_cell(st.session_state.current_sheet, selected[0], selected[1], "")
            st.session_state.selected_pos = None
            st.session_state.notice = "Đã xóa thùng khỏi vị trí."
            st.rerun()
    else:
        st.info("Bấm vào ô đang chứa thùng cần xóa.")

st.divider()
st.subheader("🗺️ Sơ đồ vị trí — bấm trực tiếp vào ô")
if action == "Di chuyển thùng" and st.session_state.move_mode:
    st.caption("Chế độ chọn đích: hãy chọn sheet đích ở Sidebar nếu cần, rồi bấm ô đích bên dưới.")

# Render every cell as a real Streamlit button. Clicking a cell selects it.
# Color is communicated with prefixes to work consistently across Streamlit versions.
current_tier = ""
for r in range(df.shape[0]):
    row_vals = [str(df.iat[r, c]).strip() for c in range(df.shape[1])]
    for val in row_vals:
        low = val.lower()
        if "tang 3" in low:
            current_tier = "Tang 3"
        elif "tang 2" in low:
            current_tier = "Tang 2"
        elif "tang 1" in low:
            current_tier = "Tang 1"

    cols = st.columns(df.shape[1], gap="small")
    for c in range(df.shape[1]):
        val = str(df.iat[r, c]).strip()
        low = val.lower()
        if "tang 3" in low:
            label = f"🟣 {val}"
        elif "tang 2" in low:
            label = f"🔵 {val}"
        elif "tang 1" in low:
            label = f"🟢 {val}"
        elif "khu" in low:
            label = f"🟠 {val}"
        elif not val:
            label = "＋ Trống"
        else:
            prefix = "🟣" if current_tier == "Tang 3" else "🔵" if current_tier == "Tang 2" else "🟢" if current_tier == "Tang 1" else "📦"
            label = f"{prefix} {val}"

        is_selected = selected == (r, c)
        if is_selected:
            label = f"👉 {val or 'Ô trống'}"
        key = f"cell_{st.session_state.current_sheet}_{r}_{c}"
        with cols[c]:
            if st.button(label, key=key, help=f"{position_text(r, c)}" + (f" · {val}" if val else "")):
                st.session_state.selected_pos = (r, c)
                st.rerun()

st.divider()
with open(EXCEL_PATH, "rb") as f:
    st.download_button(
        "📥 Tải báo cáo Excel",
        data=f.read(),
        file_name="RACK_WMS_Report.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
st.caption("Lưu ý: đóng file RACK.xlsx trong Excel trước khi thao tác để tránh xung đột khi lưu.")
